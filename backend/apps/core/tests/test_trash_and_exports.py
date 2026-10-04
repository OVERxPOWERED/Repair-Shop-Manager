import io
from datetime import timedelta

import openpyxl
import pytest
from django.core.management import call_command
from django.utils import timezone

from apps.audit.models import AuditLog
from apps.billing.models import Invoice, Payment
from apps.customers.models import Customer
from apps.devices.models import Device
from apps.jobs.models import Job

pytestmark = pytest.mark.django_db


def test_trash_list_and_permissions(world, client_for):
    """A deleted job appears only with ?deleted=true; Engineer gets 403; Owner gets 200."""
    owner_client = client_for(world.owner_a, world.shop_a)
    eng_client = client_for(world.engineer_a, world.shop_a)

    customer = Customer.objects.create(shop=world.shop_a, name="Trash Test Cust", phone="9988776655")
    device = Device.objects.create(
        shop=world.shop_a, customer=customer, category="mobile", brand_text="Apple", model="iPhone 13"
    )
    job = Job.objects.create(
        shop=world.shop_a,
        job_no=101,
        customer=customer,
        device=device,
        fault_description="Cracked screen",
        received_at=timezone.now(),
    )

    # Initially in live list
    r = owner_client.get("/api/v1/jobs/")
    assert r.status_code == 200
    job_ids = [j["id"] for j in r.json()["data"]]
    assert str(job.id) in job_ids

    # Soft delete the job
    job.soft_delete()

    # Excluded from default list
    r = owner_client.get("/api/v1/jobs/")
    assert r.status_code == 200
    job_ids = [j["id"] for j in r.json()["data"]]
    assert str(job.id) not in job_ids

    # Engineer cannot see trash (Engineer has no jobs.restore permission)
    r = eng_client.get("/api/v1/jobs/?deleted=true")
    assert r.status_code == 403

    # Owner can see trash
    r = owner_client.get("/api/v1/jobs/?deleted=true")
    assert r.status_code == 200
    trash_ids = [j["id"] for j in r.json()["data"]]
    assert str(job.id) in trash_ids


def test_restore_job_and_audit(world, client_for):
    """Restoring a deleted job brings it back to the live list and writes an audit log."""
    owner_client = client_for(world.owner_a, world.shop_a)
    customer = Customer.objects.create(shop=world.shop_a, name="Restore Cust", phone="9988776654")
    device = Device.objects.create(
        shop=world.shop_a, customer=customer, category="mobile", brand_text="Samsung", model="S22"
    )
    job = Job.objects.create(
        shop=world.shop_a,
        job_no=102,
        customer=customer,
        device=device,
        fault_description="Battery issue",
        received_at=timezone.now(),
    )
    job.soft_delete()

    r = owner_client.post(f"/api/v1/jobs/{job.id}/restore/")
    assert r.status_code == 200

    job.refresh_from_db()
    assert job.deleted_at is None

    # Check audit log
    audit = AuditLog.objects.filter(action="job.restored", entity_id=str(job.id)).first()
    assert audit is not None
    assert audit.shop == world.shop_a


def test_customer_restore_phone_conflict(world, client_for):
    """Restoring a customer when another active customer has the same phone returns 409 customer.phone_exists."""
    owner_client = client_for(world.owner_a, world.shop_a)
    phone = "9876500001"

    # Customer 1 created and soft-deleted
    c1 = Customer.objects.create(shop=world.shop_a, name="Old Customer", phone=phone)
    c1.soft_delete()

    # Customer 2 created with the same phone
    c2 = Customer.objects.create(shop=world.shop_a, name="New Customer", phone=phone)

    # Restoring c1 should fail with 409
    r = owner_client.post(f"/api/v1/customers/{c1.id}/restore/")
    assert r.status_code == 409
    assert r.json()["error"]["code"] == "customer.phone_exists"

    # But if c2 changes phone or is soft-deleted, c1 can be restored
    c2.phone = "9876500002"
    c2.save()

    r = owner_client.post(f"/api/v1/customers/{c1.id}/restore/")
    assert r.status_code == 200
    c1.refresh_from_db()
    assert c1.deleted_at is None


def test_permanent_delete_safety(world, client_for):
    """Permanent delete is rejected if financial records exist; bare records are permanently removed."""
    owner_client = client_for(world.owner_a, world.shop_a)

    customer = Customer.objects.create(shop=world.shop_a, name="Fin Customer", phone="9988001122")
    device = Device.objects.create(
        shop=world.shop_a, customer=customer, category="mobile", brand_text="Apple", model="iPad"
    )
    job = Job.objects.create(
        shop=world.shop_a,
        job_no=103,
        customer=customer,
        device=device,
        fault_description="Charging port",
        received_at=timezone.now(),
    )

    # Attach a payment to job
    import uuid

    Payment.objects.create(
        shop=world.shop_a,
        job=job,
        customer=customer,
        amount_paise=50000,
        mode="cash",
        received_by=world.owner_a,
        received_at=timezone.now(),
        idempotency_key=uuid.uuid4(),
    )
    job.soft_delete()

    # Attempt permanent delete on job with payment -> 409
    r = owner_client.delete(f"/api/v1/jobs/{job.id}/permanent/")
    assert r.status_code == 409
    assert r.json()["error"]["code"] == "trash.has_financial_records"

    # Customer also cannot be permanently deleted (has job)
    customer.soft_delete()
    r = owner_client.delete(f"/api/v1/customers/{customer.id}/permanent/")
    assert r.status_code == 409
    assert r.json()["error"]["code"] == "trash.has_financial_records"

    # Now create a bare job with no payments or invoices
    bare_cust = Customer.objects.create(shop=world.shop_a, name="Bare Cust", phone="9988009988")
    bare_device = Device.objects.create(
        shop=world.shop_a, customer=bare_cust, category="mobile", brand_text="Apple", model="iPhone 11"
    )
    bare_job = Job.objects.create(
        shop=world.shop_a,
        job_no=104,
        customer=bare_cust,
        device=bare_device,
        fault_description="Speaker",
        received_at=timezone.now(),
    )
    bare_job.soft_delete()

    r = owner_client.delete(f"/api/v1/jobs/{bare_job.id}/permanent/")
    assert r.status_code == 204
    assert not Job.all_objects.filter(id=bare_job.id).exists()

    # Verify audit row created
    audit = AuditLog.objects.filter(action="job.deleted_permanently", entity_id=str(bare_job.id)).first()
    assert audit is not None


def test_purge_trash_management_command(world, tmp_path, settings):
    """purge_trash hard-deletes records soft-deleted > 30 days ago that pass can_hard_delete."""
    old_time = timezone.now() - timedelta(days=35)
    recent_time = timezone.now() - timedelta(days=10)

    # Customer 1: old and bare -> should be purged
    c_old_bare = Customer.objects.create(shop=world.shop_a, name="Old Bare", phone="9111111111")
    c_old_bare.soft_delete()
    Customer.all_objects.filter(id=c_old_bare.id).update(deleted_at=old_time)

    # Customer 2: recent and bare -> should NOT be purged
    c_recent_bare = Customer.objects.create(shop=world.shop_a, name="Recent Bare", phone="9222222222")
    c_recent_bare.soft_delete()
    Customer.all_objects.filter(id=c_recent_bare.id).update(deleted_at=recent_time)

    # Customer 3: old with job -> should NOT be purged
    c_old_with_job = Customer.objects.create(shop=world.shop_a, name="Old With Job", phone="9333333333")
    device = Device.objects.create(
        shop=world.shop_a, customer=c_old_with_job, category="mobile", brand_text="MI", model="Redmi 10"
    )
    Job.objects.create(
        shop=world.shop_a,
        job_no=201,
        customer=c_old_with_job,
        device=device,
        fault_description="Screen",
        received_at=timezone.now(),
    )
    c_old_with_job.soft_delete()
    Customer.all_objects.filter(id=c_old_with_job.id).update(deleted_at=old_time)

    # Job: old and bare -> should be purged
    job_old_bare = Job.objects.create(
        shop=world.shop_a,
        job_no=202,
        customer=c_recent_bare,
        device=device,
        fault_description="Camera",
        received_at=timezone.now(),
    )
    job_old_bare.soft_delete()
    Job.all_objects.filter(id=job_old_bare.id).update(deleted_at=old_time)

    call_command("purge_trash")

    # c_old_bare is gone
    assert not Customer.all_objects.filter(id=c_old_bare.id).exists()
    # c_recent_bare is still in trash
    assert Customer.all_objects.filter(id=c_recent_bare.id).exists()
    # c_old_with_job is still in trash (has financial record/job)
    assert Customer.all_objects.filter(id=c_old_with_job.id).exists()
    # job_old_bare is gone
    assert not Job.all_objects.filter(id=job_old_bare.id).exists()


def test_excel_exports_permissions_and_scoping(world, client_for):
    """Front Desk gets 403; Owner gets valid .xlsx; other shop records never leak."""
    front_desk_client = client_for(world.front_desk_a, world.shop_a)
    owner_client = client_for(world.owner_a, world.shop_a)

    # Create records in Shop A
    cust_a = Customer.objects.create(shop=world.shop_a, name="Shop A Cust", phone="9876543210")
    dev_a = Device.objects.create(
        shop=world.shop_a, customer=cust_a, category="mobile", brand_text="Apple", model="iPhone 14"
    )
    job_a = Job.objects.create(
        shop=world.shop_a,
        job_no=301,
        customer=cust_a,
        device=dev_a,
        fault_description="Water damage",
        estimate_paise=150000,
        total_paise=150000,
        received_at=timezone.now(),
    )

    # Create records in Shop B
    cust_b = Customer.objects.create(shop=world.shop_b, name="Shop B Cust", phone="9123456789")
    dev_b = Device.objects.create(
        shop=world.shop_b, customer=cust_b, category="mobile", brand_text="Apple", model="iPhone 15"
    )
    job_b = Job.objects.create(
        shop=world.shop_b,
        job_no=301,
        customer=cust_b,
        device=dev_b,
        fault_description="Broken back glass",
        estimate_paise=200000,
        total_paise=200000,
        received_at=timezone.now(),
    )

    # 1. Front Desk cannot export (no data.export permission)
    r = front_desk_client.get("/api/v1/exports/jobs.xlsx")
    assert r.status_code == 403
    r = front_desk_client.get("/api/v1/exports/customers.xlsx")
    assert r.status_code == 403

    # 2. Owner can export jobs
    r = owner_client.get("/api/v1/exports/jobs.xlsx")
    assert r.status_code == 200
    assert r["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "jobs.xlsx" in r["Content-Disposition"]

    # Open with openpyxl and verify contents
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) >= 2  # header + at least 1 job row
    headers = rows[0]
    assert "Job No" in headers
    assert "Customer Phone" in headers

    # Verify Shop A job is present and unmasked phone
    phone_idx = headers.index("Customer Phone")
    phones = [row[phone_idx] for row in rows[1:]]
    assert "9876543210" in phones
    assert "9123456789" not in phones  # Shop B never leaks!

    # 3. Check audit log written
    audit = AuditLog.objects.filter(action="data.exported", entity_id="jobs", shop=world.shop_a).first()
    assert audit is not None
    assert audit.after["type"] == "jobs"

    # 4. Date range validation: > 366 days -> 400
    r = owner_client.get("/api/v1/exports/jobs.xlsx?from=2024-01-01&to=2025-02-01")
    assert r.status_code == 400

    # 5. Customers export
    r = owner_client.get("/api/v1/exports/customers.xlsx")
    assert r.status_code == 200
    wb_cust = openpyxl.load_workbook(io.BytesIO(r.content))
    ws_cust = wb_cust.active
    cust_rows = list(ws_cust.iter_rows(values_only=True))
    cust_names = [row[1] for row in cust_rows[1:]]
    assert "Shop A Cust" in cust_names
    assert "Shop B Cust" not in cust_names

    # 6. Invoices export
    inv_a = Invoice.objects.create(
        shop=world.shop_a,
        customer=cust_a,
        job=job_a,
        kind=Invoice.Kind.TAX_INVOICE,
        status=Invoice.Status.ISSUED,
        number_display="INV-001",
        issue_date=timezone.now().date(),
        subtotal_paise=150000,
        taxable_paise=150000,
        total_paise=150000,
    )
    Invoice.objects.create(
        shop=world.shop_b,
        customer=cust_b,
        job=job_b,
        kind=Invoice.Kind.TAX_INVOICE,
        status=Invoice.Status.ISSUED,
        number_display="INV-B-001",
        issue_date=timezone.now().date(),
        subtotal_paise=200000,
        taxable_paise=200000,
        total_paise=200000,
    )
    r = owner_client.get("/api/v1/exports/invoices.xlsx")
    assert r.status_code == 200
    wb_inv = openpyxl.load_workbook(io.BytesIO(r.content))
    inv_rows = list(wb_inv.active.iter_rows(values_only=True))
    inv_numbers = [row[0] for row in inv_rows[1:]]
    assert "INV-001" in inv_numbers
    assert "INV-B-001" not in inv_numbers

    # 7. Payments export
    import uuid

    Payment.objects.create(
        shop=world.shop_a,
        customer=cust_a,
        job=job_a,
        invoice=inv_a,
        amount_paise=75000,
        mode="upi",
        direction="in",
        reference="UPI12345",
        received_by=world.owner_a,
        received_at=timezone.now(),
        idempotency_key=uuid.uuid4(),
    )
    Payment.objects.create(
        shop=world.shop_b,
        customer=cust_b,
        job=job_b,
        invoice=inv_a,
        amount_paise=90000,
        mode="cash",
        direction="in",
        received_by=world.owner_b,
        received_at=timezone.now(),
        idempotency_key=uuid.uuid4(),
    )
    r = owner_client.get("/api/v1/exports/payments.xlsx")
    assert r.status_code == 200
    wb_pay = openpyxl.load_workbook(io.BytesIO(r.content))
    pay_rows = list(wb_pay.active.iter_rows(values_only=True))
    pay_refs = [row[4] for row in pay_rows[1:]]
    assert "UPI12345" in pay_refs
