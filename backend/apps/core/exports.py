import io
from datetime import date, datetime
from decimal import Decimal

from django.db import models
from django.http import HttpResponse
from openpyxl import Workbook
from rest_framework.exceptions import ValidationError

from apps.audit.services import record_audit
from apps.billing.models import Invoice, Payment
from apps.core.time import IST, today_ist
from apps.customers.models import Customer
from apps.jobs.models import Job
from apps.tenancy.viewsets import ShopScopedAPIView


def parse_date_range(request) -> tuple[date | None, date | None]:
    from_str = request.query_params.get("from")
    to_str = request.query_params.get("to")
    from_date = None
    to_date = None
    if from_str:
        try:
            from_date = datetime.strptime(from_str, "%Y-%m-%d").date()
        except ValueError as err:
            raise ValidationError("Invalid 'from' date format. Use YYYY-MM-DD.") from err
    if to_str:
        try:
            to_date = datetime.strptime(to_str, "%Y-%m-%d").date()
        except ValueError as err:
            raise ValidationError("Invalid 'to' date format. Use YYYY-MM-DD.") from err
    if from_date and to_date:
        if to_date < from_date:
            raise ValidationError("'to' date cannot be earlier than 'from' date.")
        if (to_date - from_date).days > 366:
            raise ValidationError("Date range cannot exceed 366 days.")
    elif from_date and not to_date:
        if (today_ist() - from_date).days > 366:
            raise ValidationError("Date range cannot exceed 366 days.")
    return from_date, to_date


def xlsx_response(content: bytes, filename: str) -> HttpResponse:
    response = HttpResponse(
        content,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def build_customers_workbook(queryset) -> tuple[bytes, int]:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title="Customers")
    headers = [
        "Customer ID",
        "Name",
        "Phone",
        "Alt Phone",
        "Email",
        "Address",
        "Notes",
        "Created At (IST)",
    ]
    ws.append(headers)
    count = 0
    for customer in queryset.iterator(chunk_size=2000):
        created_str = customer.created_at.astimezone(IST).strftime("%Y-%m-%d %H:%M:%S") if customer.created_at else ""
        ws.append(
            [
                str(customer.id),
                customer.name,
                customer.phone or "",
                customer.alt_phone or "",
                customer.email or "",
                customer.address or "",
                customer.notes or "",
                created_str,
            ]
        )
        count += 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), count


def build_jobs_workbook(queryset) -> tuple[bytes, int]:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title="Jobs")
    headers = [
        "Job No",
        "Customer Name",
        "Customer Phone",
        "Device",
        "IMEI/Serial",
        "Status",
        "Priority",
        "Fault Description",
        "Estimate (Rs)",
        "Total (Rs)",
        "Received At (IST)",
        "Ready At (IST)",
        "Delivered At (IST)",
    ]
    ws.append(headers)
    count = 0
    for job in (
        queryset.select_related("customer", "device").prefetch_related("device__identifiers").iterator(chunk_size=2000)
    ):
        cust_name = job.customer.name if job.customer else ""
        cust_phone = job.customer.phone if job.customer else ""
        dev_str = str(job.device) if job.device else ""
        imei = ""
        if job.device:
            idents = [i.value for i in job.device.identifiers.all() if not i.deleted_at]
            imei = idents[0] if idents else ""

        recv_str = job.received_at.astimezone(IST).strftime("%Y-%m-%d %H:%M:%S") if job.received_at else ""
        ready_str = job.ready_at.astimezone(IST).strftime("%Y-%m-%d %H:%M:%S") if job.ready_at else ""
        deliv_str = job.delivered_at.astimezone(IST).strftime("%Y-%m-%d %H:%M:%S") if job.delivered_at else ""

        ws.append(
            [
                job.job_no,
                cust_name,
                cust_phone or "",
                dev_str,
                imei,
                job.status,
                job.priority,
                job.fault_description or "",
                Decimal(job.estimate_paise) / Decimal(100),
                Decimal(job.total_paise) / Decimal(100),
                recv_str,
                ready_str,
                deliv_str,
            ]
        )
        count += 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), count


def build_invoices_workbook(queryset) -> tuple[bytes, int]:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title="Invoices")
    headers = [
        "Invoice No",
        "Kind",
        "Status",
        "Issue Date",
        "Customer Name",
        "Customer Phone",
        "Customer GSTIN",
        "Place of Supply",
        "Subtotal (Rs)",
        "Discount (Rs)",
        "Taxable (Rs)",
        "CGST (Rs)",
        "SGST (Rs)",
        "IGST (Rs)",
        "Total (Rs)",
        "Amount Paid (Rs)",
        "Balance (Rs)",
    ]
    ws.append(headers)
    count = 0
    for inv in queryset.select_related("customer").iterator(chunk_size=2000):
        cust_name = inv.customer.name if inv.customer else ""
        cust_phone = inv.customer.phone if inv.customer else ""
        ws.append(
            [
                inv.number_display or str(inv.id),
                inv.kind,
                inv.status,
                str(inv.issue_date) if inv.issue_date else "",
                cust_name,
                cust_phone or "",
                inv.customer_gstin or "",
                inv.place_of_supply_state or "",
                Decimal(inv.subtotal_paise) / Decimal(100),
                Decimal(inv.discount_paise) / Decimal(100),
                Decimal(inv.taxable_paise) / Decimal(100),
                Decimal(inv.cgst_paise) / Decimal(100),
                Decimal(inv.sgst_paise) / Decimal(100),
                Decimal(inv.igst_paise) / Decimal(100),
                Decimal(inv.total_paise) / Decimal(100),
                Decimal(inv.amount_paid_paise) / Decimal(100),
                Decimal(inv.balance_paise) / Decimal(100),
            ]
        )
        count += 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), count


def build_payments_workbook(queryset) -> tuple[bytes, int]:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title="Payments")
    headers = [
        "Payment ID",
        "Direction",
        "Mode",
        "Amount (Rs)",
        "Reference",
        "Customer Name",
        "Invoice No",
        "Job No",
        "Received At (IST)",
        "Notes",
    ]
    ws.append(headers)
    count = 0
    for pay in queryset.select_related("customer", "invoice", "job").iterator(chunk_size=2000):
        cust_name = pay.customer.name if pay.customer else ""
        inv_no = pay.invoice.number_display if pay.invoice else ""
        job_no = pay.job.job_no if pay.job else ""
        recv_str = pay.received_at.astimezone(IST).strftime("%Y-%m-%d %H:%M:%S") if pay.received_at else ""
        ws.append(
            [
                str(pay.id),
                pay.direction,
                pay.mode,
                Decimal(pay.amount_paise) / Decimal(100),
                pay.reference or "",
                cust_name,
                inv_no or "",
                job_no,
                recv_str,
                pay.notes or "",
            ]
        )
        count += 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), count


class CustomerExportView(ShopScopedAPIView):
    permission_map = {"get": "data.export"}

    def get(self, request):
        qs = Customer.objects.filter(shop=request.shop, deleted_at__isnull=True).order_by("-created_at")
        content, count = build_customers_workbook(qs)
        record_audit(
            request=request,
            action="data.exported",
            entity_type="export",
            entity_id="customers",
            after={"type": "customers", "from": None, "to": None, "rows": count},
        )
        return xlsx_response(content, "customers.xlsx")


class JobExportView(ShopScopedAPIView):
    permission_map = {"get": "data.export"}

    def get(self, request):
        from_date, to_date = parse_date_range(request)
        qs = Job.objects.filter(shop=request.shop, deleted_at__isnull=True).order_by("-received_at")
        if from_date:
            qs = qs.filter(received_at__date__gte=from_date)
        if to_date:
            qs = qs.filter(received_at__date__lte=to_date)
        content, count = build_jobs_workbook(qs)
        record_audit(
            request=request,
            action="data.exported",
            entity_type="export",
            entity_id="jobs",
            after={
                "type": "jobs",
                "from": from_date.isoformat() if from_date else None,
                "to": to_date.isoformat() if to_date else None,
                "rows": count,
            },
        )
        return xlsx_response(content, "jobs.xlsx")


class InvoiceExportView(ShopScopedAPIView):
    permission_map = {"get": "data.export"}

    def get(self, request):
        from_date, to_date = parse_date_range(request)
        qs = Invoice.objects.filter(shop=request.shop, deleted_at__isnull=True).order_by("-created_at")
        if from_date:
            qs = qs.filter(
                models.Q(issue_date__gte=from_date) | models.Q(issue_date__isnull=True, created_at__date__gte=from_date)
            )
        if to_date:
            qs = qs.filter(
                models.Q(issue_date__lte=to_date) | models.Q(issue_date__isnull=True, created_at__date__lte=to_date)
            )
        content, count = build_invoices_workbook(qs)
        record_audit(
            request=request,
            action="data.exported",
            entity_type="export",
            entity_id="invoices",
            after={
                "type": "invoices",
                "from": from_date.isoformat() if from_date else None,
                "to": to_date.isoformat() if to_date else None,
                "rows": count,
            },
        )
        return xlsx_response(content, "invoices.xlsx")


class PaymentExportView(ShopScopedAPIView):
    permission_map = {"get": "data.export"}

    def get(self, request):
        from_date, to_date = parse_date_range(request)
        qs = Payment.objects.filter(shop=request.shop).order_by("-received_at")
        if from_date:
            qs = qs.filter(received_at__date__gte=from_date)
        if to_date:
            qs = qs.filter(received_at__date__lte=to_date)
        content, count = build_payments_workbook(qs)
        record_audit(
            request=request,
            action="data.exported",
            entity_type="export",
            entity_id="payments",
            after={
                "type": "payments",
                "from": from_date.isoformat() if from_date else None,
                "to": to_date.isoformat() if to_date else None,
                "rows": count,
            },
        )
        return xlsx_response(content, "payments.xlsx")
